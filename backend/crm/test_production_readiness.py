from datetime import date
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from unittest import skipUnless
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.db import connection, connections
from django.test import override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from rest_framework.test import APIClient, APITestCase, APITransactionTestCase
from openpyxl import load_workbook

from .models import (
    AddonSale, Branch, CatalogItem, CertificateTemplate, Client, Discount, EmployeePayrollProfile,
    FinanceTransaction, GiftCertificate, GroupMembership, Lesson, MasterClass, MasterClassPayment, MasterClassSubject,
    PaymentMethod, PayrollStatement, Room, StudyGroup, Subscription, Trial, Visit,
)
from .payment_parts import sync_finance_payment_parts


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class ProductionReadinessTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user(username='audit-admin', role='admin', roles=['admin'])
        cls.branch = Branch.objects.create(name='Audit branch')
        cls.other_branch = Branch.objects.create(name='Other audit branch')
        cls.student = Client.objects.create(first_name='Audit student', branch=cls.branch)
        cls.other_student = Client.objects.create(first_name='Other student', branch=cls.branch)
        cls.cash = PaymentMethod.objects.create(name='Audit cash', code='audit-cash', is_cash=True)
        cls.card = PaymentMethod.objects.create(name='Audit card', code='audit-card')
        cls.subject = MasterClassSubject.objects.create(name='Audit master class')

    def setUp(self):
        self.client.force_authenticate(self.admin)

    def income(self, split=True):
        item = FinanceTransaction.objects.create(
            transaction_type='income', amount=100, source='manual', branch=self.branch,
        )
        parts = [{'payment_method': self.cash, 'amount': Decimal('40')} ,
                 {'payment_method': self.card, 'amount': Decimal('60')}]
        sync_finance_payment_parts(item, parts if split else None, legacy_payment_method=self.cash)
        return item

    def subscription(self, **kwargs):
        values = dict(client=self.student, branch=self.branch, title='Audit course',
                      start_date=date(2026, 9, 1), total_visits=4, remaining_visits=4, price=100)
        values.update(kwargs)
        return Subscription.objects.create(**values)

    def master_class(self):
        item = MasterClass.objects.create(subject=self.subject, title=self.subject.name,
                                         branch=self.branch, starts_at=timezone.now(), price=100)
        item.participants.add(self.student)
        return item

    def test_rejected_finance_amount_patch_is_atomic(self):
        item = self.income()
        response = self.client.patch(f'/api/finance/{item.id}/', {'amount': '120.00'}, format='json')
        self.assertEqual(response.status_code, 400)
        item.refresh_from_db()
        self.assertEqual(item.amount, Decimal('100'))
        self.assertEqual(sum(part.amount for part in item.payment_parts.all()), item.amount)

    def test_invalid_payment_part_amounts_return_400_without_writes(self):
        self.client.raise_request_exception = False
        for amount in ('abc', 'NaN', 'Infinity', [], {}, '1e1000'):
            with self.subTest(amount=amount):
                response = self.client.post('/api/finance/', {
                    'transaction_type': 'income', 'amount': '100',
                    'payment_parts': [{'payment_method': self.cash.id, 'amount': amount}],
                }, format='json')
                self.assertEqual(response.status_code, 400)
        self.assertFalse(FinanceTransaction.objects.exists())

    def test_split_subscription_status_patch_preserves_payment(self):
        item = self.income()
        subscription = self.subscription(paid_amount=100, finance_transaction=item)
        response = self.client.patch(f'/api/subscriptions/{subscription.id}/', {'status': 'paused'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        item.refresh_from_db()
        self.assertEqual(item.payment_parts.count(), 2)
        self.assertEqual(item.amount, Decimal('100'))

    def test_split_trial_stage_patch_preserves_payment(self):
        item = self.income()
        trial = Trial.objects.create(client=self.student, branch=self.branch, scheduled_at=timezone.now(),
                                     price=100, payment_date=date(2026, 9, 1), finance_transaction=item)
        response = self.client.patch(f'/api/trials/{trial.id}/', {'stage': 'attended'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(item.payment_parts.count(), 2)

    def test_master_class_split_payment_comment_patch_preserves_parts(self):
        mc = self.master_class()
        item = self.income()
        payment = MasterClassPayment.objects.create(master_class=mc, amount=100,
                                                    payment_date=date(2026, 9, 1), finance_transaction=item)
        response = self.client.patch(f'/api/master-classes/{mc.id}/payments/{payment.id}/',
                                     {'comment': 'Updated'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(item.payment_parts.count(), 2)

    def test_master_class_payment_method_switch_uses_selected_method(self):
        mc = self.master_class()
        item = self.income(split=False)
        payment = MasterClassPayment.objects.create(master_class=mc, amount=100,
                                                    payment_date=date(2026, 9, 1), finance_transaction=item)
        response = self.client.patch(f'/api/master-classes/{mc.id}/payments/{payment.id}/',
                                     {'payment_method': self.card.id}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        item.refresh_from_db()
        self.assertEqual(item.payment_method_id, self.card.id)
        self.assertEqual(item.payment_parts.get().payment_method_id, self.card.id)
        self.assertEqual(response.data['payment_parts'][0]['payment_method'], self.card.id)

    def test_master_class_payment_missing_parent_returns_404(self):
        self.client.raise_request_exception = False
        response = self.client.post('/api/master-classes/999999/payments/',
                                    {'amount': 10, 'payment_method': self.cash.id}, format='json')
        self.assertEqual(response.status_code, 404)

    def test_master_class_payment_respects_branch_filter(self):
        mc = self.master_class()
        response = self.client.post(f'/api/master-classes/{mc.id}/payments/?branch={self.other_branch.id}',
                                    {'amount': 10, 'payment_method': self.cash.id}, format='json')
        self.assertEqual(response.status_code, 404)
        self.assertFalse(mc.payments.exists())

    def test_sale_comment_does_not_reprice_or_mark_partial_sale_paid(self):
        product = CatalogItem.objects.create(category='product', name='Audit product', price=100, owner_name='Original')
        response = self.client.post('/api/addon-sales/', {
            'branch': self.branch.id, 'items': [{'catalog_item': product.id, 'quantity': 1}],
            'payment_amount': '40', 'payment_method': self.cash.id,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        sale = AddonSale.objects.get(pk=response.data['id'])
        product.price = 200
        product.owner_name = 'Changed'
        product.save()
        response = self.client.patch(f'/api/addon-sales/{sale.id}/', {'comment': 'Updated'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        sale.refresh_from_db()
        self.assertEqual(sale.total_price, Decimal('100'))
        self.assertEqual(sale.payment_amount, Decimal('40'))
        self.assertEqual(sale.finance_transaction.amount, Decimal('40'))
        self.assertEqual(sale.items.get().owner_name, 'Original')

    def test_subscription_explicit_lesson_count_sets_matching_balance(self):
        service = CatalogItem.objects.create(category='service', name='Four lessons', price=100, lessons_count=4)
        response = self.client.post('/api/subscriptions/', {
            'client': self.student.id, 'service': service.id, 'total_visits': 8,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['remaining_visits'], 8)

    def test_subscription_status_patch_preserves_inactive_discount_snapshot(self):
        discount = Discount.objects.create(name='Historical', discount_type='percentage', value=10, is_active=False)
        subscription = self.subscription(discount=discount, discount_name='Original',
                                         discount_type='percentage', discount_value=5, discount_amount=5)
        response = self.client.patch(f'/api/subscriptions/{subscription.id}/', {'status': 'paused'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        subscription.refresh_from_db()
        self.assertEqual(subscription.discount_amount, Decimal('5'))
        self.assertEqual(subscription.discount_name, 'Original')

    def test_visit_delete_restores_deducted_lesson(self):
        subscription = self.subscription(remaining_visits=3)
        visit = Visit.objects.create(client=self.student, subscription=subscription, visited_at=timezone.now(),
                                     status='attended', lesson_deducted=True)
        response = self.client.delete(f'/api/visits/{visit.id}/')
        self.assertEqual(response.status_code, 204)
        subscription.refresh_from_db()
        self.assertEqual(subscription.remaining_visits, 4)

    def test_visit_rejects_other_client_subscription(self):
        subscription = self.subscription(client=self.other_student)
        response = self.client.post('/api/visits/', {
            'client': self.student.id, 'subscription': subscription.id,
            'visited_at': timezone.now().isoformat(), 'status': 'attended',
        }, format='json')
        self.assertEqual(response.status_code, 400)
        subscription.refresh_from_db()
        self.assertEqual(subscription.remaining_visits, 4)

    def test_visit_deduction_flag_cannot_be_overridden(self):
        subscription = self.subscription(remaining_visits=3)
        visit = Visit.objects.create(client=self.student, subscription=subscription, visited_at=timezone.now(),
                                     status='attended', lesson_deducted=True)
        response = self.client.patch(f'/api/visits/{visit.id}/', {'lesson_deducted': False}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        subscription.refresh_from_db()
        self.assertEqual(subscription.remaining_visits, 3)

    def test_visit_date_is_business_date(self):
        visit = Visit.objects.create(client=self.student, visited_at='2026-09-13T20:30:00Z')
        response = self.client.get(f'/api/visits/{visit.id}/')
        self.assertEqual(str(response.data['date']), '2026-09-14')

    def test_visit_subscription_switch_restores_original_and_deducts_new(self):
        original = self.subscription(remaining_visits=3)
        replacement = self.subscription()
        visit = Visit.objects.create(client=self.student, subscription=original, visited_at=timezone.now(),
                                     status='attended', lesson_deducted=True)
        response = self.client.patch(f'/api/visits/{visit.id}/', {'subscription': replacement.id}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        original.refresh_from_db()
        replacement.refresh_from_db()
        self.assertEqual(original.remaining_visits, 4)
        self.assertEqual(replacement.remaining_visits, 3)

    def test_linked_finance_delete_does_not_leave_fictitious_payment(self):
        item = self.income(split=False)
        subscription = self.subscription(paid_amount=100, finance_transaction=item)
        response = self.client.delete(f'/api/finance/{item.id}/')
        self.assertEqual(response.status_code, 409)
        subscription.refresh_from_db()
        self.assertEqual(subscription.finance_transaction_id, item.id)

    def test_linked_finance_amount_cannot_diverge_from_subscription(self):
        item = self.income(split=False)
        self.subscription(paid_amount=100, finance_transaction=item)
        response = self.client.patch(f'/api/finance/{item.id}/', {'amount': '75'}, format='json')
        self.assertEqual(response.status_code, 400)
        item.refresh_from_db()
        self.assertEqual(item.amount, Decimal('100'))

    def test_paid_subscription_delete_preserves_history(self):
        item = self.income(split=False)
        subscription = self.subscription(paid_amount=100, finance_transaction=item)
        response = self.client.delete(f'/api/subscriptions/{subscription.id}/')
        self.assertEqual(response.status_code, 409)
        self.assertTrue(Subscription.objects.filter(pk=subscription.id).exists())

    def test_linked_finance_full_form_can_update_comment_without_changing_payment(self):
        item = self.income()
        self.subscription(paid_amount=100, finance_transaction=item)
        payload = self.client.get(f'/api/finance/{item.id}/').data
        payload['comment'] = 'Updated comment'
        response = self.client.patch(f'/api/finance/{item.id}/', payload, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        item.refresh_from_db()
        self.assertEqual(item.comment, 'Updated comment')
        self.assertEqual(item.amount, Decimal('100'))
        self.assertEqual(item.payment_parts.count(), 2)

    def test_paid_payroll_cannot_be_approved_again(self):
        statement = PayrollStatement.objects.create(employee=self.admin, date_from=date(2026, 9, 1),
                                                    date_to=date(2026, 9, 30), status='paid')
        response = self.client.post(f'/api/payroll/{statement.id}/approve/', {}, format='json')
        self.assertEqual(response.status_code, 400)
        statement.refresh_from_db()
        self.assertEqual(statement.status, 'paid')

    def test_payroll_adjustment_updates_total_immediately(self):
        EmployeePayrollProfile.objects.create(employee=self.admin, pay_type='monthly', monthly_salary=1000)
        statement = PayrollStatement.objects.create(employee=self.admin, date_from=date(2026, 9, 1),
                                                    date_to=date(2026, 9, 30), base_amount=1000, total_amount=1000)
        response = self.client.patch(f'/api/payroll/{statement.id}/', {'manual_adjustment': '50'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Decimal(response.data['total_amount']), Decimal('1050'))

    def test_master_class_legacy_status_patch_keeps_title(self):
        mc = self.master_class()
        mc.subject = None
        mc.save()
        response = self.client.patch(f'/api/master-classes/{mc.id}/', {'stage': 'booked'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['title'], self.subject.name)

    def test_master_class_unrelated_patch_keeps_title_snapshot(self):
        mc = self.master_class()
        self.subject.name = 'Renamed subject'
        self.subject.save()
        response = self.client.patch(f'/api/master-classes/{mc.id}/', {'stage': 'booked'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['title'], mc.title)

    def test_certificate_batch_edit_does_not_reduce_shared_payment(self):
        template = CertificateTemplate.objects.create(name='Audit certificate', amount_type='fixed', fixed_amount=100,
                                                       sale_discount_percent=0, validity_days=365)
        response = self.client.post('/api/certificates/bulk-create/', {
            'template': template.id, 'quantity': 3, 'face_value': 100, 'purchaser_client': self.student.id,
            'payment_parts': [{'payment_method': self.cash.id, 'amount': 300}],
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        certificate = response.data['certificates'][0]
        item = FinanceTransaction.objects.get(pk=certificate['finance_transaction'])
        response = self.client.patch(f"/api/certificates/{certificate['id']}/", {'recipient_name': 'Recipient'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        item.refresh_from_db()
        self.assertEqual(item.amount, Decimal('300'))
        self.assertEqual(item.payment_parts.get().amount, Decimal('300'))

    def test_certificate_cannot_change_nominal_after_issue(self):
        template = CertificateTemplate.objects.create(name='Range certificate', amount_type='range', min_amount=1,
                                                       sale_discount_percent=0, validity_days=365)
        response = self.client.post('/api/certificates/', {
            'template': template.id, 'face_value': 100, 'purchaser_client': self.student.id,
            'payment_parts': [{'payment_method': self.cash.id, 'amount': 100}],
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        certificate = response.data
        response = self.client.patch(f"/api/certificates/{certificate['id']}/", {'face_value': '200'}, format='json')
        self.assertEqual(response.status_code, 400)

    def test_export_visit_datetime_is_business_timezone(self):
        Visit.objects.create(client=self.student, visited_at='2026-09-13T20:30:00Z')
        response = self.client.get('/api/export/visits/')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.cell(2, 2).value, '2026-09-14 01:30')

    def test_trial_conversion_invalid_values_return_400(self):
        self.client.raise_request_exception = False
        for overrides in ({'price': 'NaN'}, {'price': '-10'}, {'total_visits': 'oops'}, {'payment_amount': '-10'},
                          {'start_date': '2026-02-30'}, {'end_date': '2020-01-01'}):
            with self.subTest(overrides=overrides):
                trial = Trial.objects.create(client=self.student, branch=self.branch, scheduled_at=timezone.now())
                response = self.client.post(f'/api/trials/{trial.id}/convert-to-subscription/', {
                    'title': 'Course', 'start_date': '2026-09-01', 'total_visits': 4,
                    'price': '100', 'payment_amount': '0', **overrides,
                }, format='json')
                self.assertEqual(response.status_code, 400)
        self.assertFalse(Subscription.objects.exists())

    def test_payroll_payment_failure_rolls_back_transaction(self):
        statement = PayrollStatement.objects.create(employee=self.admin, date_from=date(2026, 9, 1),
                                                    date_to=date(2026, 9, 30), status='approved', total_amount=100)
        self.client.raise_request_exception = False
        with patch('crm.views.sync_finance_payment_parts', side_effect=RuntimeError('Simulated storage failure')):
            response = self.client.post(f'/api/payroll/{statement.id}/mark-paid/', {
                'payment_parts': [{'payment_method': self.cash.id, 'amount': 100}],
            }, format='json')
        self.assertEqual(response.status_code, 500)
        self.assertFalse(FinanceTransaction.objects.exists())
        statement.refresh_from_db()
        self.assertEqual(statement.status, 'approved')

    def test_master_class_initial_payment_invalid_amount_returns_400(self):
        self.client.raise_request_exception = False
        response = self.client.post('/api/master-classes/', {
            'subject': self.subject.id, 'client': self.student.id, 'branch': self.branch.id,
            'starts_at': timezone.now().isoformat(), 'price': 100,
            'initial_payment': {'amount': 'NaN', 'payment_method': self.cash.id},
        }, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(MasterClass.objects.exists())

    def lesson(self):
        group = StudyGroup.objects.create(name='Audit group', branch=self.branch)
        GroupMembership.objects.create(client=self.student, group=group)
        return Lesson.objects.create(group=group, branch=self.branch, lesson_date=date(2026, 9, 14),
                                     start_time='16:00', end_time='17:00')

    def test_attendance_invalid_second_row_rolls_back_first_row(self):
        lesson = self.lesson()
        subscription = self.subscription()
        response = self.client.post(f'/api/lessons/{lesson.id}/attendance/', {'items': [
            {'client': self.student.id, 'subscription': subscription.id, 'status': 'attended'},
            {'client': self.other_student.id, 'status': 'attended'},
        ]}, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Visit.objects.exists())
        subscription.refresh_from_db()
        self.assertEqual(subscription.remaining_visits, 4)

    def test_attendance_rejects_foreign_subscription(self):
        lesson = self.lesson()
        subscription = self.subscription(client=self.other_student)
        response = self.client.post(f'/api/lessons/{lesson.id}/attendance/', {'items': [
            {'client': self.student.id, 'subscription': subscription.id, 'status': 'attended'},
        ]}, format='json')
        self.assertEqual(response.status_code, 400)
        subscription.refresh_from_db()
        self.assertEqual(subscription.remaining_visits, 4)

    def test_lesson_rejects_reversed_time_and_foreign_room(self):
        room = Room.objects.create(name='Other room', branch=self.other_branch)
        for overrides in ({'end_time': '15:00'}, {'room': room.id}):
            with self.subTest(overrides=overrides):
                response = self.client.post('/api/lessons/', {
                    'branch': self.branch.id, 'lesson_date': '2026-09-14', 'start_time': '16:00', 'end_time': '17:00',
                    **overrides,
                }, format='json')
                self.assertEqual(response.status_code, 400)
        self.assertFalse(Lesson.objects.exists())

    def test_certificate_redemption_rechecks_locked_balance(self):
        from .views import GiftCertificateViewSet
        certificate = GiftCertificate.objects.create(code='AUDIT', face_value=100, sale_price=100,
                                                     remaining_amount=100, issued_at=date(2026, 9, 1),
                                                     valid_until=date(2027, 9, 1))
        GiftCertificate.objects.filter(pk=certificate.pk).update(remaining_amount=40, status='partially_used')
        with patch.object(GiftCertificateViewSet, 'get_object', return_value=certificate):
            response = self.client.post(f'/api/certificates/{certificate.id}/redeem/', {'amount': 50}, format='json')
        self.assertEqual(response.status_code, 400)
        certificate.refresh_from_db()
        self.assertEqual(certificate.remaining_amount, Decimal('40'))
        self.assertFalse(certificate.redemptions.exists())

    def test_master_class_list_queries_do_not_grow_per_client(self):
        self.master_class()
        with CaptureQueriesContext(connection) as first:
            response = self.client.get('/api/master-classes/')
        self.assertEqual(response.status_code, 200)
        for _ in range(5):
            self.master_class()
        with CaptureQueriesContext(connection) as larger:
            response = self.client.get('/api/master-classes/')
        self.assertEqual(response.status_code, 200)
        self.assertLessEqual(len(larger), len(first) + 2, f'Queries grew from {len(first)} to {len(larger)}')


@skipUnless(connection.vendor == 'postgresql', 'Requires PostgreSQL row locks')
@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class ProductionConcurrencyTests(APITransactionTestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user(username='concurrency-admin', role='admin', roles=['admin'])
        self.cash = PaymentMethod.objects.create(name='Concurrency cash', code='concurrency-cash', is_cash=True)

    def concurrent_posts(self, view_class, url, payload):
        barrier = Barrier(2, timeout=15)
        original_get_object = view_class.get_object

        def synchronized_get_object(view):
            instance = original_get_object(view)
            barrier.wait()
            return instance

        def send(_):
            try:
                client = APIClient()
                client.force_authenticate(get_user_model().objects.get(pk=self.admin.pk))
                response = client.post(url, payload, format='json')
                return response.status_code
            finally:
                connections.close_all()

        # Both requests first read the old state; correctness must come from the locked re-read.
        with patch.object(view_class, 'get_object', synchronized_get_object):
            with ThreadPoolExecutor(max_workers=2) as pool:
                return sorted(pool.map(send, range(2)))

    def test_concurrent_payroll_payment_creates_one_expense(self):
        from .views import PayrollStatementViewSet
        statement = PayrollStatement.objects.create(employee=self.admin, date_from=date(2026, 9, 1),
                                                    date_to=date(2026, 9, 30), status='approved', total_amount=100)
        statuses = self.concurrent_posts(PayrollStatementViewSet, f'/api/payroll/{statement.id}/mark-paid/', {
            'payment_parts': [{'payment_method': self.cash.id, 'amount': 100}],
        })
        self.assertEqual(statuses, [200, 200])
        self.assertEqual(FinanceTransaction.objects.count(), 1)
        statement.refresh_from_db()
        self.assertEqual(statement.status, 'paid')
        self.assertEqual(statement.finance_transaction.amount, Decimal('100'))

    def test_concurrent_trial_conversion_creates_one_subscription(self):
        from .views import TrialViewSet
        student = Client.objects.create(first_name='Concurrency student')
        trial = Trial.objects.create(client=student, scheduled_at=timezone.now())
        statuses = self.concurrent_posts(TrialViewSet, f'/api/trials/{trial.id}/convert-to-subscription/', {
            'title': 'Course', 'start_date': '2026-09-01', 'total_visits': 4, 'price': 100,
            'payment_amount': 100, 'payment_method': self.cash.id,
        })
        self.assertEqual(statuses, [201, 400])
        self.assertEqual(Subscription.objects.count(), 1)
        self.assertEqual(FinanceTransaction.objects.count(), 1)

    def test_concurrent_redemptions_cannot_overspend_certificate(self):
        from .views import GiftCertificateViewSet
        certificate = GiftCertificate.objects.create(code='CONCURRENT', face_value=100, sale_price=100,
                                                     remaining_amount=100, issued_at=date(2026, 9, 1),
                                                     valid_until=date(2099, 9, 1))
        statuses = self.concurrent_posts(GiftCertificateViewSet, f'/api/certificates/{certificate.id}/redeem/', {'amount': 60})
        self.assertEqual(statuses, [200, 400])
        certificate.refresh_from_db()
        self.assertEqual(certificate.remaining_amount, Decimal('40'))
        self.assertEqual(certificate.redemptions.count(), 1)
